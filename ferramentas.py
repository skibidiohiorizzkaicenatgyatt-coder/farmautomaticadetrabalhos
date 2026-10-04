import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONTE = "Times New Roman"
_MARCACAO = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _campo(run, instrucao):
    for tipo, texto in [("begin", None), (None, instrucao), ("end", None)]:
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def _configurar_estilos(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    normal.font.size = Pt(12)
    for nome, tam in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12), ("Title", 20), ("Subtitle", 14)]:
        st = doc.styles[nome]
        st.font.name = FONTE
        st.element.rPr.rFonts.set(qn("w:ascii"), FONTE)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), FONTE)
        st.font.size = Pt(tam)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = nome != "Subtitle"
        st.font.italic = nome == "Heading 3"
        st.paragraph_format.space_before = Pt(18 if nome == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True


def _texto(p, texto):
    """Escreve o texto no parágrafo, entendendo **negrito** e *itálico*."""
    for parte in _MARCACAO.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            p.add_run(parte[2:-2]).bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            p.add_run(parte[1:-1]).italic = True
        else:
            p.add_run(parte)


def _paragrafo(doc, texto):
    for par in texto.split("\n\n"):
        par = par.strip()
        if not par:
            continue
        p = doc.add_paragraph()
        _texto(p, par)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)


def _topico(doc, texto):
    p = doc.add_paragraph(style="List Bullet")
    _texto(p, texto)
    p.paragraph_format.line_spacing = 1.5


def _numerado(doc, texto, n):
    p = doc.add_paragraph()
    _texto(p, f"{n}. {texto}")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.63)
    p.paragraph_format.line_spacing = 1.5


def _recuado(doc, texto):
    """Bloco recuado à esquerda, fonte menor, espaçamento simples (ementa, citação longa)."""
    p = doc.add_paragraph()
    _texto(p, texto)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(7)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    for r in p.runs:
        r.font.size = Pt(11)


def _referencia(doc, texto):
    p = doc.add_paragraph()
    _texto(p, texto)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)


def _numero_pagina(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _campo(p.add_run(), "PAGE")


def criar_docx(trabalho: dict, caminho: str) -> str:
    """Monta o DOCX. O visual é fixo; a estrutura vem da lista 'blocos' que a IA devolve."""
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(3); sec.left_margin = Cm(3)
        sec.bottom_margin = Cm(2); sec.right_margin = Cm(2)
    _configurar_estilos(doc)
    _numero_pagina(doc)

    # capa
    if trabalho.get("incluir_capa", True):
        for _ in range(6):
            doc.add_paragraph()
        t = doc.add_paragraph(trabalho["titulo"], style="Title")
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if trabalho.get("subtitulo"):
            s = doc.add_paragraph(trabalho["subtitulo"], style="Subtitle")
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
    elif trabalho.get("incluir_titulo", True):
        doc.add_paragraph(trabalho["titulo"], style="Title")

    # sumário
    if trabalho.get("incluir_sumario", True):
        doc.add_heading("Sumário", level=1)
        _campo(doc.add_paragraph().add_run(), 'TOC \\o "1-2" \\h \\z \\u')
        doc.add_page_break()

    # corpo: um bloco por vez
    blocos = trabalho.get("blocos", [])
    n_numerado = 0
    i = 0
    while i < len(blocos):
        tipo = blocos[i].get("tipo", "paragrafo")
        texto = (blocos[i].get("texto") or "").strip()

        if tipo == "referencia":  # referências seguidas ficam em ordem alfabética
            j = i
            while j < len(blocos) and blocos[j].get("tipo") == "referencia":
                j += 1
            refs = [(b.get("texto") or "").strip() for b in blocos[i:j]]
            for ref in sorted([r for r in refs if r], key=str.lower):
                _referencia(doc, ref)
            i, n_numerado = j, 0
            continue

        i += 1
        if tipo != "numerado":
            n_numerado = 0
        if tipo == "quebra_pagina":
            doc.add_page_break()
            continue
        if not texto:
            continue

        if tipo in ("titulo1", "titulo2", "titulo3"):
            doc.add_heading(texto, level=int(tipo[-1]))
        elif tipo == "topico":
            _topico(doc, texto)
        elif tipo == "numerado":
            n_numerado += 1
            _numerado(doc, texto, n_numerado)
        elif tipo == "recuado":
            _recuado(doc, texto)
        else:  # "paragrafo" e qualquer tipo desconhecido
            _paragrafo(doc, texto)

    doc.save(caminho)
    return caminho