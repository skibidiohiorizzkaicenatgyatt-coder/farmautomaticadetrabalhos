import shutil
from pathlib import Path

import visao

LIMITE_POR_ARQUIVO = 25000   # instruções e modelos: caracteres por arquivo
LIMITE_MODELOS = 30000       # modelos: caracteres no total
LIMITE_TOTAL = 60000         # modelos + trechos das fontes que vão para a IA
LIMITE_LEITURA = 2_000_000   # fontes: caracteres lidos por arquivo (segurança)
PAPEIS = ["instrucoes", "modelos", "fontes"]


def _ler_pdf(caminho):
    """Texto do PDF, página por página. Páginas sem texto (escaneadas) são lidas pelo Claude."""
    from pypdf import PdfReader
    nome = Path(caminho).name
    paginas = [(p.extract_text() or "").strip() for p in PdfReader(str(caminho)).pages]
    vazias = [i for i, t in enumerate(paginas, 1) if len(t) < 30]
    if vazias and visao.OCR_ATIVO:
        alvo = vazias[:visao.MAX_PAGINAS_OCR]
        if len(vazias) > len(alvo):
            print(f"   AVISO: {len(vazias)} páginas sem texto em {nome}; lendo só as {len(alvo)} primeiras "
                  "(MAX_PAGINAS_OCR no visao.py)", flush=True)
        print(f"   lendo {len(alvo)} página(s) escaneada(s) de {nome} com o Claude...", flush=True)
        try:
            for n, t in visao.transcrever_paginas_pdf(caminho, alvo).items():
                paginas[n - 1] = "" if t.strip() == "[sem texto legível]" else t
        except Exception as e:
            print(f"   AVISO: não consegui ler as páginas escaneadas ({type(e).__name__}: {e})", flush=True)
    return "\n".join(f"[p. {i}]\n{t}" for i, t in enumerate(paginas, 1) if t)

def _ler_docx(caminho):
    from docx import Document
    d = Document(caminho)
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for tabela in d.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(c.text.strip() for c in linha.cells))
    return "\n".join(partes)

def _ler_xlsx(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True, read_only=True)
    partes = []
    for ws in wb.worksheets:
        partes.append(f"[Planilha: {ws.title}]")
        for linha in ws.iter_rows(values_only=True):
            if any(c is not None for c in linha):
                partes.append(" | ".join("" if c is None else str(c) for c in linha))
    return "\n".join(partes)

def _ler_pptx(caminho):
    from pptx import Presentation
    partes = []
    for i, slide in enumerate(Presentation(caminho).slides, 1):
        partes.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                partes.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            partes.append("Notas: " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(partes)

def _ler_texto(caminho):
    return Path(caminho).read_text(encoding="utf-8", errors="ignore")

def _ler_imagem(caminho):
    if not visao.OCR_ATIVO:
        raise RuntimeError("leitura de imagens desativada (OCR_ATIVO no visao.py)")
    print(f"   lendo a imagem {Path(caminho).name} com o Claude...", flush=True)
    return visao.transcrever_imagem(caminho)

LEITORES = {
    ".pdf": _ler_pdf, ".docx": _ler_docx, ".xlsx": _ler_xlsx, ".pptx": _ler_pptx,
    ".txt": _ler_texto, ".md": _ler_texto, ".csv": _ler_texto,
    ".png": _ler_imagem, ".jpg": _ler_imagem, ".jpeg": _ler_imagem,
    ".webp": _ler_imagem, ".gif": _ler_imagem,
}


def ler_entrada(pasta="entrada"):
    """Lê entrada/instrucoes, entrada/modelos e entrada/fontes (cria se não existirem).
    Retorna (instrucoes, modelos, docs): dois textos e a lista de fontes [{'nome', 'texto'}].
    As fontes NÃO são cortadas aqui: se forem grandes, os trechos relevantes são escolhidos depois."""
    raiz = Path(pasta)
    for papel in PAPEIS:
        (raiz / papel).mkdir(parents=True, exist_ok=True)

    soltos = [a.name for a in raiz.iterdir() if a.is_file()]
    if soltos:
        print(f"   AVISO: arquivos soltos em '{pasta}' são ignorados, coloque nas subpastas: {', '.join(soltos)}", flush=True)

    textos, docs, total_modelos = {"instrucoes": "", "modelos": ""}, [], 0
    for papel in PAPEIS:
        blocos = []
        for arq in sorted((raiz / papel).iterdir()):
            if not arq.is_file() or arq.name.startswith("~$"):
                continue
            leitor = LEITORES.get(arq.suffix.lower())
            if not leitor:
                print(f"   (ignorado, formato não suportado: {papel}/{arq.name})", flush=True)
                continue
            try:
                texto = leitor(arq).strip()
            except Exception as e:
                print(f"   (não consegui ler {papel}/{arq.name}: {type(e).__name__}: {e})", flush=True)
                continue
            if not texto:
                print(f"   (sem texto extraível: {papel}/{arq.name})", flush=True)
                continue

            if papel == "fontes":
                if len(texto) > LIMITE_LEITURA:
                    texto = texto[:LIMITE_LEITURA]
                    print(f"   (arquivo enorme, só os primeiros {LIMITE_LEITURA} caracteres foram lidos: {arq.name})", flush=True)
                docs.append({"nome": arq.name, "texto": texto})
                print(f"   lido: fontes/{arq.name} ({len(texto)} caracteres)", flush=True)
                continue

            if len(texto) > LIMITE_POR_ARQUIVO:
                texto = texto[:LIMITE_POR_ARQUIVO] + "\n[... arquivo cortado por tamanho ...]"
                print(f"   (cortado em {LIMITE_POR_ARQUIVO} caracteres: {papel}/{arq.name})", flush=True)
            if papel == "modelos":
                if total_modelos + len(texto) > LIMITE_MODELOS:
                    print(f"   (limite de modelos atingido, ignorado: {arq.name})", flush=True)
                    continue
                total_modelos += len(texto)
            blocos.append(f"=== {arq.name} ===\n{texto}")
            print(f"   lido: {papel}/{arq.name} ({len(texto)} caracteres)", flush=True)
        if papel != "fontes":
            textos[papel] = "\n\n".join(blocos)

    return textos["instrucoes"], textos["modelos"], docs


def arquivar_entrada(pasta_entrada, destino):
    """Move instrucoes/, modelos/ e fontes/ pra dentro de 'destino' e recria as pastas vazias.
    Nunca derruba o programa: se não conseguir mover (ex: arquivo aberto no Word), copia."""
    raiz, destino = Path(pasta_entrada), Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    for papel in PAPEIS:
        origem = raiz / papel
        try:
            if origem.exists():
                shutil.move(str(origem), str(destino / papel))
        except Exception as e:
            print(f"   AVISO: não consegui mover '{papel}' ({type(e).__name__}). Copiando.", flush=True)
            try:
                shutil.copytree(origem, destino / papel, dirs_exist_ok=True)
                shutil.rmtree(origem, ignore_errors=True)
            except Exception:
                print(f"   AVISO: '{papel}' ficou na entrada. Feche os arquivos abertos e mova manualmente.", flush=True)
        (raiz / papel).mkdir(parents=True, exist_ok=True)